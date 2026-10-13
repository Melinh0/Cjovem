from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import pyperclip
import time
import re
import os
from datetime import datetime

from config import MOODLE_URL, PLANILHAS_DIR, require_moodle_credentials

NOME_TURMA = "TQ14"
PROFESSOR_TURMA = "Lucas Vieira"
MENTOR_RESPONSAVEL = "Yago Melo da Costa"
COORDENADOR_CELULA = "Ariel"
AULA_SEMANA = "2.0"

USERNAME, PASSWORD = require_moodle_credentials()

options = webdriver.ChromeOptions()
options.add_argument('--start-maximized')
driver = webdriver.Chrome(options=options)

try:
    driver.get(MOODLE_URL)
    time.sleep(2)

    cpf_input = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located((By.ID, "inputName"))
    )
    cpf_input.send_keys(USERNAME)
    time.sleep(2)

    password_input = driver.find_element(By.ID, "inputPassword")
    password_input.send_keys(PASSWORD)
    password_input.send_keys(Keys.RETURN)
    time.sleep(2)

    course_link = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located((By.XPATH, "//span[contains(text(), '[CICLO 2] CIÊNCIA DE DADOS - MÓDULO INTERMEDIÁRIO')]"))
    )
    driver.execute_script("arguments[0].click();", course_link)
    time.sleep(2)

    participants_link = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located((By.XPATH, "//a[contains(text(), 'Participantes')]"))
    )
    driver.execute_script("arguments[0].click();", participants_link)
    time.sleep(2)

    show_all_button = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located((By.XPATH, "//a[@data-action='showcount']"))
    )
    driver.execute_script("arguments[0].click();", show_all_button)
    time.sleep(15)

    body = driver.find_element(By.TAG_NAME, "body")
    body.send_keys(Keys.CONTROL + 'a')
    body.send_keys(Keys.CONTROL + 'c')
    page_content = pyperclip.paste()

    lines = page_content.split('\n')
    alunos = []
    for line in lines:
        parts = line.split('\t')
        if len(parts) == 5:
            nome = parts[0]
            nome = re.sub(r'^[A-ZÀ-Ý]{1,2}(?=[A-ZÁÉÍÓÚÂÊÔÃÕÄËÏÖÜÇ])', '', nome).strip()
            alunos.append({
                'nome': nome,
                'email': parts[1],
                'ultimo_acesso': parts[4]
            })

    alunos = [aluno for aluno in alunos if aluno['nome'] not in [
        "Lucas Vieira Alves",
        "Yago Melo da Costa"
    ]]

    alunos.sort(key=lambda x: x['nome'])

    data_atual = datetime.now()
    data_relatorio = data_atual.strftime("%d/%m/%Y")
    data_mensagem = data_atual.strftime("%d/%m")

    wb = Workbook()
    ws = wb.active
    ws.title = f"{NOME_TURMA} - Relatório"

    ws.merge_cells('B3:H3')
    ws['B3'] = "RELATÓRIO DE AÇÕES DA MENTORIA - CAPACITA BRASIL CJOVEM -CIÊNCIA DE DADOS - NÍVEL BÁSICO"
    ws['B3'].font = Font(bold=True, size=14)
    ws['B3'].alignment = Alignment(horizontal='center')

    ws['B4'] = "NOME DA TURMA:"
    ws['C4'] = NOME_TURMA
    ws['E4'] = "DATA DO PRESENTE RELATÓRIO:"
    ws['F4'] = data_relatorio
    ws['G4'] = "PROFESSOR(A) DA TURMA:"
    ws['H4'] = PROFESSOR_TURMA

    ws['B5'] = "MENTOR(A) RESPONSÁVEL:"
    ws['C5'] = MENTOR_RESPONSAVEL
    ws['E5'] = "COORDENADOR DE CÉLULA RESPONSÁVEL:"
    ws['F5'] = COORDENADOR_CELULA
    ws['G5'] = "AULA DA SEMANA:"
    ws['H5'] = AULA_SEMANA

    cabecalhos = [
        "NOME DO ALUNO(A)", "EMAIL", "STATUS NO CURSO",
        "ÚLTIMO ACESSO NO AVA", "STATUS NO DISCORD",
        "DESCRIÇÃO DA ÚLTIMA AÇÃO", "RETORNO DO ALUNO"
    ]

    for col, cabecalho in enumerate(cabecalhos, start=2):
        cell = ws.cell(row=7, column=col, value=cabecalho)
        cell.font = Font(bold=True)

    for i, aluno in enumerate(alunos, start=8):
        status_curso = "Cursando" if "Nunca" not in aluno['ultimo_acesso'] else "Não iniciado"
        status_discord = "Entrou no discord" if status_curso == "Cursando" else "Não entrou no discord"
        retorno = "RETORNO" if status_curso == "Cursando" else "SEM RETORNO"

        ws.cell(row=i, column=2, value=aluno['nome'])
        ws.cell(row=i, column=3, value=aluno['email'])
        ws.cell(row=i, column=4, value=status_curso)
        ws.cell(row=i, column=5, value=aluno['ultimo_acesso'])
        ws.cell(row=i, column=6, value=status_discord)
        ws.cell(row=i, column=7, value=f"Enviado mensagem dia {data_mensagem} para envio de material e aviso de aula.")
        ws.cell(row=i, column=8, value=retorno)

    ws_acoes = wb.create_sheet("Ações diárias - Mentores - Chec")
    acoes = [
        "Diariamente",
        "Mensagem no grupo do whatspp e discord",
        "lembrete das aulas síncronas",
        "lembrete dos prazos das aulas (semanais)",
        "lembrete \"tá precisando de alguma coisa?\"",
        "lembrete \"assine o termo de desistência\"",
        "enviar notícias no off-topic"
    ]

    for i, acao in enumerate(acoes, start=1):
        ws_acoes.cell(row=i, column=1, value=acao)

    nome_arquivo = f"Status_Alunos_{NOME_TURMA}.xlsx"
    PLANILHAS_DIR.mkdir(parents=True, exist_ok=True)
    caminho_completo = PLANILHAS_DIR / nome_arquivo
    wb.save(caminho_completo)
    print(f"Planilha salva em: {os.path.abspath(caminho_completo)}")

except Exception as e:
    print(f"Erro: {e}")

finally:
    driver.quit()
